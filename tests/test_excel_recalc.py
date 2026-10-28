"""Independent check: LibreOffice recalculates every formula and must agree with the cached values we wrote."""
import io
import shutil
import subprocess

import openpyxl
import pytest

from app.models import AnalysisOptions
from tests.conftest import make_png, wait_for

pytestmark = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


def test_cached_values_match_recalculation(app_service, tmp_path):
    aid = app_service.create_analysis("s.png", make_png(), AnalysisOptions(review_before_search=False))
    wait_for(app_service, aid, {"completed"})
    data = app_service.export_xlsx(aid)
    src = tmp_path / "in.xlsx"
    src.write_bytes(data)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    a = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    b = openpyxl.load_workbook(out / "in.xlsx", data_only=True)
    errors = {"#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "Err:502", "Err:508"}
    checked = 0
    for ws in a.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v2 = b[ws.title][c.coordinate].value
                assert str(v2) not in errors, (ws.title, c.coordinate, v2)
                if isinstance(c.value, (int, float)) and isinstance(v2, (int, float)):
                    assert abs(c.value - v2) < 0.01, (ws.title, c.coordinate, c.value, v2)
                    checked += 1
                elif c.value not in (None, "") and v2 not in (None, ""):
                    assert str(c.value) == str(v2), (ws.title, c.coordinate, c.value, v2)
    assert checked > 50
